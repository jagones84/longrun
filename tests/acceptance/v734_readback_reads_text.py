#!/usr/bin/env python3
"""v734 — the read-back reads the WORDS of a deliverable, not only its shape.

EV-37. On JobBench `online_merchants/task1` the workbook really held a bar chart
(`xl/charts/chart1.xml`, live `$A$2:$A$6` refs) and the grader still wrote "no
chart, no conditional formatting": it reads TEXT. The same blindness cost three
other criteria — the report never wrote the required filename, the catalog mean
was printed as 4.1 instead of 4.28, the sample size (178 units) was missing.

So the read-back now extracts the text of every deliverable and states two
mechanical facts the agent can act on:

  A  a deliverable's TEXT is readable for every format we ship (xlsx, pdf, csv,
     md/txt, docx via stdlib zip)
  B  an INPUT whose filename appears in no deliverable is named as such
  C  a chart is announced as a picture — its numbers must also be in cells
  D  an unreadable file is a fact, never a crash

Deterministic and offline: temp dirs, no model, no network.
"""
import os
import sys
import tempfile
import time
import zipfile

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
TMP = tempfile.mkdtemp(prefix="longrun-v734-")
sys.path.insert(0, os.path.join(REPO, "src"))

from longrun.plan import artifacts as A  # noqa: E402

results = []


def check(name, ok, detail=""):
    results.append(bool(ok))
    print(("PASS " if ok else "FAIL ") + name
          + ((" :: " + str(detail)[:300]) if detail and not ok else ""))


WS = os.path.join(TMP, "ws")
GIVEN = os.path.join(WS, "task_folder")
OUT = os.path.join(WS, "output")
for d in (GIVEN, OUT):
    os.makedirs(d, exist_ok=True)

OLD = time.time() - 3600
for name, body in (("product_performance_Q4.csv", "SKU,Return_Rate\nSKU-007,8.2%\n"),
                   ("customer_inquiries.csv", "Inquiry_ID,Subject\nINQ-011,Colour\n")):
    p = os.path.join(GIVEN, name)
    with open(p, "w", encoding="utf-8") as f:
        f.write(body)
    os.utime(p, (OLD, OLD))

SINCE = OLD + 60

# ---- A: text of every format ------------------------------------------------
try:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference

    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    ws.append(["SKU", "Return Rate"])
    ws.append(["SKU-007", 8.2])
    wb.create_sheet("Executive Summary").append(["Catalog mean return 4.28%"])
    ch = BarChart()
    ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    ws.add_chart(ch, "D2")
    xlsx = os.path.join(OUT, "Report.xlsx")
    wb.save(xlsx)
except Exception as exc:  # noqa: BLE001
    print("SKIP A: openpyxl unavailable (%s)" % exc)


def _write(name, body):
    p = os.path.join(OUT, name)
    with open(p, "w", encoding="utf-8") as f:
        f.write(body)
    return p


txt = _write("notes.md", "Catalog mean return 4.28% over 17 products.")
docx = os.path.join(OUT, "notice.docx")
with zipfile.ZipFile(docx, "w") as z:
    z.writestr("[Content_Types].xml", "<Types/>")
    z.writestr("word/document.xml",
               "<w:document><w:body><w:p><w:r><w:t>Dear supplier: 178 units</w:t>"
               "</w:r></w:p></w:body></w:document>")

if os.path.exists(xlsx):
    check("A1 an xlsx yields its CELL text (including the chart's own table)",
          "SKU-007" in A.deliverable_text(xlsx)
          and "Catalog mean return 4.28" in A.deliverable_text(xlsx),
          A.deliverable_text(xlsx)[:200])
check("A2 a markdown/text deliverable yields its words",
      "4.28" in A.deliverable_text(txt))
check("A3 a docx yields its words with the stdlib only",
      "178 units" in A.deliverable_text(docx), A.deliverable_text(docx)[:120])
check("A4 an unknown/binary suffix yields '' instead of raising",
      A.deliverable_text(os.path.join(OUT, "nope.xyz")) == "")

# ---- B: a source named / not named -----------------------------------------
_t = A.review(WS, "produce the report", since=SINCE) or ""
check("B1 with no source named, both inputs are called out",
      "NOT NAMED in any deliverable" in _t and "product_performance_Q4.csv" in _t
      and "customer_inquiries.csv" in _t, _t[-500:])

_write("Report.xlsx.txt", "")   # no-op keeper so the next write is newest
if os.path.exists(xlsx):
    wb2 = __import__("openpyxl").load_workbook(xlsx)
    wb2["Executive Summary"].append(
        ["Source: product_performance_Q4.csv and customer_inquiries.csv"])
    wb2.save(xlsx)
else:
    _write("Report.xlsx", "")
_t2 = A.review(WS, "produce the report", since=SINCE) or ""
check("B2 once the deliverables NAME both files, nothing is called out",
      "NOT NAMED in any deliverable" not in _t2, _t2[-500:])

# ---- C: a chart is a picture ------------------------------------------------
check("C1 a chart triggers the 'restate its numbers' line",
      ("a chart is a PICTURE" in _t2) if os.path.exists(xlsx) else True,
      _t2[-300:])

# ---- D: a broken file ------------------------------------------------------
with open(os.path.join(OUT, "broken.xlsx"), "wb") as f:
    f.write(b"not a zip")
check("D1 an unreadable workbook is a fact, not a crash",
      A.deliverable_text(os.path.join(OUT, "broken.xlsx")) == ""
      and (A.review(WS, "x", since=SINCE) or "") != "")

# ---- E: a data URL is never rendered as a page -------------------------------
from longrun.tools import tools as T  # noqa: E402


def _fake(url, body):
    """Point urllib at a canned body (offline)."""
    class _R:
        def read(self, _n=None):
            return body

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    T.urllib.request.urlopen = lambda req, timeout=0: _R()


_orig_open = T.urllib.request.urlopen
_fake("x", b"<!DOCTYPE html><html>page</html>")
_data, _err = T._web_fetch("https://github.com/o/r/blob/master/apparel.csv", 4096)
check("E1 a .csv URL that answers with HTML is reported, not rendered",
      _data is None and "HTML PAGE" in (_err or ""), _err)
check("E2 the answer names the RAW url to ask for",
      "raw.githubusercontent.com/o/r/master/apparel.csv" in (_err or ""), _err)

_fake("x", b"handle,title\ndark-denim-top,Dark Denim Top ripped hem effect\n")
_data2, _err2 = T._web_fetch("https://example.org/data/apparel.csv", 4096)
check("E3 a real data file arrives byte-identical (no page rendering)",
      _err2 is None and "ripped hem effect" in (_data2 or {}).get("text", ""),
      (_data2, _err2))
T.urllib.request.urlopen = _orig_open

print("\n---")
print("v734: %d/%d PASS" % (sum(results), len(results)))
sys.exit(0 if all(results) else 1)
