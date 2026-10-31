#!/usr/bin/env python3
"""v731 — the harness READS BACK the artifacts a turn wrote before it closes.

Found 2026-10-30 on JobBench `online_merchants/task1`: Longrun delivered both
files in 54 s and scored 6/58, while Harness² — whose loop keeps re-reading its
own output for 50-100 steps — scored 48/58. The judge's failures were things the
harness could have SHOWN it: "Sheet 3 has no visualization", "the catalog
averages are not quoted". A status-only done-gate (`_finish_gated`) cannot see a
blank sheet; a mechanical read-back can.

  A  the inventory describes what the turn WROTE (sheets, ranges, charts, columns)
  B  inputs, source files and older files are never reported as deliverables
  C  the read-back is a CHECKLIST with the request in front of the model
  D  the gate fires ONCE per turn, and not at all when nothing was written
  E  it is wired into every place a final answer is accepted

Deterministic and offline: a temporary workspace, no model, no network.
"""
import os
import sys
import tempfile
import time

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
TMP = tempfile.mkdtemp(prefix="longrun-v731-")
sys.path.insert(0, os.path.join(REPO, "src"))

from longrun.plan import artifacts as A  # noqa: E402

results = []


def check(name, ok, detail=""):
    results.append(bool(ok))
    print(("PASS " if ok else "FAIL ") + name
          + ((" :: " + str(detail)[:400]) if detail and not ok else ""))


WS = os.path.join(TMP, "ws")
OUT = os.path.join(WS, "output")
os.makedirs(OUT, exist_ok=True)
os.makedirs(os.path.join(WS, "task_folder"), exist_ok=True)

# an input staged BEFORE the turn (must never be reported)
with open(os.path.join(WS, "task_folder", "input.csv"), "w", encoding="utf-8") as f:
    f.write("a,b\n1,2\n")
_old = time.time() - 3600
os.utime(os.path.join(WS, "task_folder", "input.csv"), (_old, _old))

SINCE = time.time() - 1

# what the turn wrote: a workbook with a chart on its second sheet
from openpyxl import Workbook  # noqa: E402
from openpyxl.chart import BarChart, Reference  # noqa: E402

wb = Workbook()
data = wb.active
data.title = "Data"
for row in (("sku", "returns"), ("SKU-007", 8.2), ("SKU-016", 7.1)):
    data.append(row)
viz = wb.create_sheet("Viz")
viz.append(("sku", "returns"))
viz.append(("SKU-007", 8.2))
chart = BarChart()
chart.add_data(Reference(viz, min_col=2, min_row=1, max_row=2), titles_from_data=True)
viz.add_chart(chart, "D2")
wb.save(os.path.join(OUT, "report.xlsx"))

with open(os.path.join(OUT, "table.csv"), "w", encoding="utf-8") as f:
    f.write("inquiry,verdict\nINQ-001,isolated\nINQ-003,systemic\n")
with open(os.path.join(OUT, "notes.md"), "w", encoding="utf-8") as f:
    f.write("# notes\nthree findings\n")
with open(os.path.join(OUT, "helper.py"), "w", encoding="utf-8") as f:
    f.write("print('not a deliverable')\n")
with open(os.path.join(OUT, "empty.txt"), "w", encoding="utf-8") as f:
    f.write("")

# A ----------------------------------------------------------------------
inv = {e["path"]: e for e in A.inventory(WS, since=SINCE)}
check("A1 the workbook the turn wrote is inventoried",
      "output/report.xlsx" in inv, sorted(inv))
xf = (inv.get("output/report.xlsx") or {}).get("facts") or {}
check("A2 its sheets are named with their used range",
      [s["name"] for s in xf.get("sheets", [])] == ["Data", "Viz"]
      and xf["sheets"][0]["rows"] == 3 and xf["sheets"][0]["cols"] == 2, xf)
check("A3 a chart is COUNTED (the fact a status gate cannot see)",
      xf.get("charts") == 1, xf)
check("A4 a csv is described by header and row count",
      (inv.get("output/table.csv") or {}).get("facts", {}).get("columns") ==
      ["inquiry", "verdict"]
      and (inv.get("output/table.csv") or {}).get("facts", {}).get("rows") == 2,
      (inv.get("output/table.csv") or {}).get("facts"))
check("A5 an empty artifact is reported as EMPTY",
      "EMPTY" in A._line(inv["output/empty.txt"]), A._line(inv["output/empty.txt"]))

# A7/A8 — FORMATTING is not content. Measured 2026-10-31 on JobBench
# `medical_secretaries/task1`: the Referral Log reported `20x30` to the read-back
# and held ONE filled row (the other 17 referrals were styled empty cells), so the
# agent read "20 rows" and closed the turn on a workbook with one row of data.
from openpyxl.styles import PatternFill  # noqa: E402

_wb2 = Workbook()
_hdr = _wb2.active
_hdr.title = "Header Only"
_hdr.append(("Referral_ID", "Patient"))
for _r in range(2, 21):
    for _c in range(1, 31):
        _hdr.cell(row=_r, column=_c).fill = PatternFill("solid", fgColor="FFEB9C")
_one = _wb2.create_sheet("One Row")
_one.append(("Referral_ID", "Patient"))
_one.append(("REF-2026-0218", "Thompson"))
for _r in range(3, 21):
    for _c in range(1, 31):
        _one.cell(row=_r, column=_c).fill = PatternFill("solid", fgColor="FFEB9C")
_wb2.save(os.path.join(OUT, "hollow.xlsx"))
_hf = {e["path"]: e for e in A.inventory(WS, since=SINCE)}["output/hollow.xlsx"]["facts"]
check("A7 a styled-but-empty sheet reports how many rows are FILLED",
      [s.get("filled_rows") for s in _hf["sheets"]] == [1, 2] and _hf["sheets"][0]["rows"] >= 20,
      _hf["sheets"])
_line_h = A._line({e["path"]: e for e in A.inventory(WS, since=SINCE)}["output/hollow.xlsx"])
check("A8 a sheet whose content stops at the header is flagged NO DATA",
      "NO DATA in Header Only" in _line_h and "[2 row(s) filled]" in _line_h, _line_h)
check("A9 a fully filled sheet prints no filled-range qualifier",
      "filled]" not in A._line(inv["output/report.xlsx"]), A._line(inv["output/report.xlsx"]))

# B ----------------------------------------------------------------------
check("B1 the staged INPUT is not a deliverable of this turn",
      "task_folder/input.csv" not in inv and not any(
          e["path"].startswith("task_folder") for e in inv.values()), sorted(inv))
check("B2 source code is not audited as an artifact",
      "output/helper.py" not in inv, sorted(inv))
_state = os.path.join(WS, "state_dir")
os.makedirs(_state, exist_ok=True)
with open(os.path.join(_state, "harness_note.md"), "w", encoding="utf-8") as f:
    f.write("a note the HARNESS wrote about the turn\n")
check("A6 the harness's own state inside the workspace is never an artifact",
      "state_dir/harness_note.md" in {e["path"] for e in A.inventory(WS, since=SINCE)}
      and "state_dir/harness_note.md" not in {
          e["path"] for e in A.inventory(WS, since=SINCE, exclude=[_state])},
      "exclude=")
_before = A.inventory(WS, since=time.time() + 60)
check("B3 files older than the turn are excluded", _before == [], _before)

# C ----------------------------------------------------------------------
text = A.review(WS, "produce a workbook with a chart and a supplier notice",
                since=SINCE) or ""
check("C1 the read-back names each artifact with its facts",
      "READ-BACK" in text and "output/report.xlsx" in text and "charts=1" in text, text[:300])
check("C2 it forwards the request so the model can check coverage",
      "produce a workbook with a chart" in text, text[-300:])
check("C3 it is a checklist, not a grade (no score, an instruction to fix)",
      "NOT delivered" in text and "fix the artifact now" in text, text[-300:])
check("C4 nothing written → no read-back at all", A.review(TMP + "/nowhere", "", since=0) is None
      and A.review(os.path.join(TMP, "empty-dir"), "", since=0) is None)

# F — the files the turn was GIVEN are accounted for -----------------------
for _name, _body in (("alpha.csv", "x,y\n1,2\n"), ("beta.json", "{}"),
                     ("gamma.txt", "notes\n")):
    _p = os.path.join(WS, "task_folder", _name)
    with open(_p, "w", encoding="utf-8") as f:
        f.write(_body)
    os.utime(_p, (_old, _old))

_given = A.inputs(WS, since=SINCE)
check("F1 the inputs the turn was GIVEN are listed",
      {"task_folder/alpha.csv", "task_folder/beta.json",
       "task_folder/gamma.txt"} <= set(_given), _given)
check("F2 a file the turn WROTE is never reported as an input",
      "output/report.xlsx" not in _given and "output/table.csv" not in _given, _given)
_text_in = A.review(WS, "x", since=SINCE) or ""
check("F3 the read-back demands an account of every input",
      "INPUTS you were GIVEN" in _text_in and "explicitly dismissed" in _text_in,
      _text_in[-400:])

# G — a closed step with no SOURCE is a claim, and it is named -------------
os.environ.setdefault("LONGRUN_GRAPH_DIR", os.path.join(TMP, "graph"))
os.makedirs(os.environ["LONGRUN_GRAPH_DIR"], exist_ok=True)
from longrun.plan import taskgraph as TG  # noqa: E402

_g = TG.ensure("v731-sess", session_id="v731-sess")
TG.begin_plan(_g)
TG.add_node(_g, "read the inputs", node_id="n1")
TG.update_node(_g, "n1", status="done", evidence="read 3 tables",
               sources=["task_folder/alpha.csv:1"])
TG.add_node(_g, "compare the averages", node_id="n2")
TG.update_node(_g, "n2", status="done", evidence="the averages match")
TG.add_node(_g, "write the notice", node_id="n3")
TG.update_node(_g, "n3", status="superseded", reason="not needed: the client asked "
               "for the workbook only")
TG.save(_g)
_claims = TG.claim_only_nodes(TG.load("v731-sess"))
check("G1 a done step backed by a source is proof, not a claim", "n1" not in _claims, _claims)
check("G2 a done step with only the model's word is reported, once",
      _claims == ["n2"], _claims)
check("G3 a superseded step is not a claim to fix", "n3" not in _claims, _claims)

# D ----------------------------------------------------------------------
_state = {"left": 1}
_first = A.review(WS, "x", since=SINCE)
check("D1 a written turn produces exactly one read-back", bool(_first) is True)


def _gate(state, ws=WS, since=SINCE):
    """The same once-per-turn rule the agent applies, without a live session."""
    if int(state.get("left") or 0) <= 0:
        return None
    t = A.review(ws, "x", since=since)
    if not t:
        state["left"] = 0
        return None
    state["left"] -= 1
    return t


check("D2 the gate consumes its budget: the second call is silent", _gate(_state) and _gate(_state) is None)
check("D3 an empty turn burns the budget instead of retrying",
      _gate({"left": 1}, ws=os.path.join(TMP, "empty-dir")) is None)

# D4/D5 — an incomplete delivery earns ONE more cycle, a clean one earns none.
_clean = os.path.join(TMP, "clean")
os.makedirs(_clean, exist_ok=True)
with open(os.path.join(_clean, "ok.csv"), "w", encoding="utf-8") as f:
    f.write("a,b\n1,2\n")
check("D4 a mechanically incomplete delivery is reported as a gap",
      any("hollow" in g for g in A.gaps(WS, since=SINCE))
      and any("empty.txt" in g for g in A.gaps(WS, since=SINCE)), A.gaps(WS, since=SINCE))
check("D5 a clean delivery has no gap, so it earns no extra cycle",
      A.gaps(_clean, since=0) == [], A.gaps(_clean, since=0))

# D6-D8 — the REQUEST can make a fact a missing deliverable. Measured 2026-10-31:
# the read-back printed `conditional-formatting=0`, the model answered "done", and
# the rubric weighted 10 for red/yellow highlighting scored 0.
import shutil  # noqa: E402

_vw = os.path.join(TMP, "visual")
os.makedirs(_vw, exist_ok=True)
shutil.copy(os.path.join(OUT, "hollow.xlsx"), os.path.join(_vw, "book.xlsx"))
_vis = A.gaps(_vw, request="produce a workbook with a bar chart of returns", since=0)
check("D6 a requested chart is a GAP when no workbook has one",
      any("charts=0" in g for g in _vis), _vis)
_fmt = A.gaps(_vw, request="highlight urgent rows with conditional formatting", since=0)
check("D7 a requested formatting rule is a GAP when every workbook reports 0",
      any("conditional-formatting=0" in g for g in _fmt), _fmt)
_plain = A.gaps(_vw, request="summarise the referral table", since=0)
check("D8 a request that names no visual adds no visual gap",
      not any("charts=0" in g or "conditional-formatting=0" in g for g in _plain), _plain)
_wchart = A.gaps(WS, request="produce a workbook with a bar chart", since=SINCE)
check("D9 a workbook that HAS the chart satisfies the request",
      not any("charts=0" in g for g in _wchart), _wchart)

# E ----------------------------------------------------------------------
SRC = open(os.path.join(REPO, "src", "longrun", "agent", "agent.py"), encoding="utf-8").read()
check("E1 every floor that accepts a final answer runs the read-back first",
      SRC.count("_drv = _deliverable_review_text(") >= 4, SRC.count("_drv = _deliverable_review_text("))
check("E2 it is resolved on the session's OWN workspace",
      "_chat_workspace(sess)" in SRC.split("def _deliverable_review_text")[1][:1400],
      "workspace resolver")
check("E3 the audit can never break a turn (guarded import + except)",
      "except Exception:  # noqa: BLE001" in SRC.split("def _deliverable_review_text")[1][:2600])
check("E4 it is fed as a harness turn the operator can see",
      SRC.count('_inject(_drv, "review")') >= 4, SRC.count('_inject(_drv, "review")'))
check("E5 a gap earns one extra cycle; the wiring is the one that raises `left`",
      "_art2.gaps(" in SRC and 'state["left"] = max(int(state["left"]), 1)' in SRC,
      "gaps wiring")
check("E6 a stop that leaves NOTHING written earns one named delivery round",
      "DELIVERY FLOOR" in SRC and "_delivery_floor_used = False" in SRC
      and "def _stop_floor_text(" in SRC, "delivery floor")
check("E8 the floor is granted from BOTH stop paths, through ONE helper",
      # The two work loops (budget exhaustion and the plain stop) share the
      # helper: a fix in one cannot miss the other (regression measured
      # 2026-10-31 — the read-back is spent by then, so only the floor fires).
      SRC.count("_stop_floor_text(sess, _deliv_since, _delivery_floor_used,") == 2
      and SRC.count("_delivery_floor_used = True") == 2
      and SRC.count('_inject(_floor, "continue")') == 2,
      "helper call sites")
check("E9 the floor needs OPEN work and a FAILED stop, not a clean close",
      # A `goal_reached` / `user_stop` close (or zero open steps) is not a
      # delivery that went missing: granting the round there added a model call
      # after the final answer (v510 regression).
      'reason not in ("budget", "no_progress")' in SRC
      and "if not open_count or reason not in" in SRC, "floor guard")
check("E7 zero artifacts is SAID OUT LOUD when the request names a file",
      "written NO file" in (A.nothing_written("produce a workbook of referrals") or "")
      and A.nothing_written("how many sessions are running?") is None,
      A.nothing_written("produce a workbook of referrals"))

print("\n---")
print("v731: %d/%d PASS" % (sum(results), len(results)))
sys.exit(0 if all(results) else 1)
