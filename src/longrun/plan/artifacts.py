#!/usr/bin/env python3
"""Read-back of the ARTIFACTS a turn wrote — mechanical, task-agnostic.

Why this module exists
----------------------
A turn can end with the files it was asked for and still answer nothing: an empty
sheet, a table where a chart was required, a PDF with one line, a workbook whose
third sheet is blank. Measured 2026-10-30 on JobBench `online_merchants/task1`:
Longrun delivered both artifacts in 54 s and scored 6/58 — the judge's failures
were things the harness could have SHOWN it ("Sheet 3 has no visualization",
"the catalog averages are not quoted"), while Harness² — whose loop keeps
re-reading its own output for 50-100 steps — scored 48/58.

The harness cannot judge meaning, so it does what it CAN do mechanically: it
opens every document the turn produced and reports its structure back to the
agent before the turn is allowed to close. Nothing here knows what a rubric is,
what the task is, or which fields matter; it only describes the files.

Design rules
------------
* Pure and offline: no model call, no network, no server import.
* Only DOCUMENT artifacts are audited (see `DELIVERABLE_EXTS`): a source-file
  edit is not a deliverable, and auditing every `.py` would be noise.
* Every reader degrades to a size/error line when its library is missing — the
  audit can make a turn SLOWER, never break it.
* The result is facts, never a verdict: the agent decides what to fix.
"""
import csv
import os
import re
import time
from pathlib import Path

#: Files a task can be graded on. Source code is deliberately absent.
DELIVERABLE_EXTS = (".xlsx", ".xlsm", ".csv", ".pdf", ".docx", ".pptx",
                    ".md", ".txt", ".html", ".htm")

#: Directories that are inputs, caches or tooling — never "what I wrote".
#: Tool/cache directories only. An INPUT folder is deliberately NOT skipped: the
#: mtime rule already separates "given to the turn" from "written by it", and a
#: module that hard-codes a benchmark's folder name is not generic.
SKIP_DIRS = {".git", "__pycache__", "node_modules", ".venv", "venv",
             ".pytest_cache", ".mypy_cache"}

MAX_FILES = 8
MAX_CSV_ROWS = 5000
MAX_TEXT_CHARS = 200_000
#: A big workbook/PDF is described by its size only: opening a 200 MB xlsx to
#: count charts would make the audit the slowest part of the turn.
MAX_WORKBOOK_BYTES = 25 * 1024 * 1024
MAX_PDF_BYTES = 40 * 1024 * 1024


def _mtime(path):
    try:
        return os.path.getmtime(path)
    except OSError:
        return 0.0


#: Row cap for the FILLED-range count. A sheet whose cells were styled (borders,
#: conditional formatting, column widths) advertises a used range far larger than
#: its content, and walking every cell of a 100k-row workbook inside a read-back
#: would cost more than the fact is worth.
MAX_SCAN_ROWS = 5000


def _filled_span(ws, limit=MAX_SCAN_ROWS):
    """(rows, cols) that actually HOLD something: a blank styled cell is not data."""
    rows = cols = 0
    try:
        for i, values in enumerate(ws.iter_rows(values_only=True)):
            if i >= limit:
                break
            width = 0
            for j, value in enumerate(values, 1):
                if value is not None and str(value).strip():
                    width = j
            if width:
                rows += 1
                cols = max(cols, width)
    except Exception:  # noqa: BLE001 — an unreadable cell is a fact, not a crash
        pass
    return rows, cols


def _workbook_facts(path):
    """Sheets with their used range AND their filled range, plus charts/formatting.

    `max_row`/`max_column` are what openpyxl calls the used range: they grow with
    FORMATTING, not with data. Measured 2026-10-31 on JobBench
    `medical_secretaries/task1`: the arm's Referral Log reported `20x30` and held
    exactly ONE filled row (the rest was styled empty), the read-back said `20x30`,
    the agent trusted it and closed the turn. The read-back must therefore state
    how much of the range is actually CONTENT.
    """
    try:
        from openpyxl import load_workbook
    except Exception:  # noqa: BLE001 — a missing reader is a fact, not a crash
        return {"error": "openpyxl missing"}
    try:
        wb = load_workbook(path)
    except Exception as exc:  # noqa: BLE001
        return {"error": "unreadable: %s" % type(exc).__name__}
    sheets = []
    charts = cf = 0
    for ws in wb.worksheets:
        charts += len(getattr(ws, "_charts", []) or [])
        try:
            cf += len(list(ws.conditional_formatting))
        except Exception:  # noqa: BLE001
            pass
        filled_rows, filled_cols = _filled_span(ws)
        sheets.append({"name": ws.title, "rows": int(ws.max_row or 0),
                       "cols": int(ws.max_column or 0),
                       "filled_rows": filled_rows,
                       "filled_cols": filled_cols})
    try:
        wb.close()
    except Exception:  # noqa: BLE001
        pass
    return {"sheets": sheets, "charts": charts, "conditional_formatting": cf}


def _csv_facts(path):
    try:
        with open(path, encoding="utf-8", errors="replace", newline="") as f:
            reader = csv.reader(f)
            header = next(reader, []) or []
            rows = 0
            for _ in reader:
                rows += 1
                if rows >= MAX_CSV_ROWS:
                    break
    except OSError as exc:
        return {"error": str(exc)[:120]}
    return {"columns": [str(c)[:40] for c in header[:12]], "rows": rows}


def _pdf_facts(path):
    try:
        from pypdf import PdfReader
    except Exception:  # noqa: BLE001
        try:
            from PyPDF2 import PdfReader  # type: ignore
        except Exception:  # noqa: BLE001
            return {"error": "pypdf missing"}
    try:
        reader = PdfReader(path)
        pages = len(reader.pages)
        text = ""
        if pages:
            text = reader.pages[0].extract_text() or ""
    except Exception as exc:  # noqa: BLE001
        return {"error": "unreadable: %s" % type(exc).__name__}
    return {"pages": pages, "page1_chars": len(text.strip())}


def _text_facts(path):
    try:
        data = Path(path).read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        return {"error": str(exc)[:120]}
    return {"lines": data.count("\n") + 1, "chars": len(data),
            "empty": not data.strip(),
            "head": " ".join(data.split())[:160]}


def _xlsx_text(path, limit=MAX_TEXT_CHARS):
    """Every cell value as text — what a reader WITHOUT a renderer can see."""
    try:
        from openpyxl import load_workbook
    except Exception:  # noqa: BLE001
        return ""
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:  # noqa: BLE001
        return ""
    out, used = [], 0
    try:
        for ws in wb.worksheets:
            out.append(str(ws.title))
            for row in ws.iter_rows(values_only=True):
                line = " ".join(str(c) for c in row if c is not None)
                if not line:
                    continue
                out.append(line)
                used += len(line) + 1
                if used >= limit:
                    return "\n".join(out)[:limit]
    except Exception:  # noqa: BLE001 — one broken sheet must not sink the audit
        pass
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass
    return "\n".join(out)[:limit]


def _pdf_text(path, pages=10, limit=MAX_TEXT_CHARS):
    """The first `pages` pages of a PDF as text (empty when no reader)."""
    try:
        from pypdf import PdfReader
    except Exception:  # noqa: BLE001
        try:
            from PyPDF2 import PdfReader  # type: ignore
        except Exception:  # noqa: BLE001
            return ""
    try:
        reader = PdfReader(path)
        parts = [(p.extract_text() or "") for p in reader.pages[:pages]]
        return "\n".join(parts)[:limit]
    except Exception:  # noqa: BLE001
        return ""


def _ooxml_text(path, prefix, limit=MAX_TEXT_CHARS):
    """Text of an OOXML part set (docx/pptx) with the stdlib only."""
    import re
    import zipfile
    try:
        with zipfile.ZipFile(path) as z:
            names = sorted(n for n in z.namelist()
                           if n.startswith(prefix) and n.endswith(".xml"))
            blob, used = [], 0
            for name in names:
                data = z.read(name).decode("utf-8", "replace")
                hits = re.findall(r"<a:t>([^<]*)</a:t>|<w:t[^>]*>([^<]*)</w:t>", data)
                text = " ".join(a or b for a, b in hits)
                blob.append(text)
                used += len(text)
                if used >= limit:
                    break
            return "\n".join(blob)[:limit]
    except Exception:  # noqa: BLE001
        return ""


def deliverable_text(path, limit=MAX_TEXT_CHARS):
    """The words a reader can search INSIDE a deliverable (capped, never raises).

    EV-37: the read-back used to count structure (`charts=1`) and never read the
    words. A rubric — or a colleague — credits a figure by FINDING it, so the
    question "is the source named, is the number written down" is mechanical and
    worth answering.
    """
    suffix = Path(path).suffix.lower()
    try:
        if suffix in (".xlsx", ".xlsm"):
            if os.path.getsize(path) > MAX_WORKBOOK_BYTES:
                return ""
            return _xlsx_text(path, limit)
        if suffix == ".pdf":
            if os.path.getsize(path) > MAX_PDF_BYTES:
                return ""
            return _pdf_text(path, limit=limit)
        if suffix == ".docx":
            return _ooxml_text(path, "word/", limit)
        if suffix == ".pptx":
            return _ooxml_text(path, "ppt/slides/", limit)
        if suffix in (".csv", ".md", ".txt", ".html", ".htm"):
            return Path(path).read_text(encoding="utf-8", errors="replace")[:limit]
    except Exception:  # noqa: BLE001 — the audit must never break a turn
        return ""
    return ""


def facts(path):
    """Structural facts for one file (never raises)."""
    suffix = Path(path).suffix.lower()
    size = 0
    try:
        size = os.path.getsize(path)
    except OSError:
        pass
    out = {"size": size}
    if size == 0:
        out["empty"] = True
    if suffix in (".xlsx", ".xlsm"):
        if size > MAX_WORKBOOK_BYTES:
            out["skipped"] = "workbook over %d MB" % (MAX_WORKBOOK_BYTES // (1024 * 1024))
        else:
            out.update(_workbook_facts(path))
    elif suffix == ".csv":
        out.update(_csv_facts(path))
    elif suffix == ".pdf":
        if size > MAX_PDF_BYTES:
            out["skipped"] = "pdf over %d MB" % (MAX_PDF_BYTES // (1024 * 1024))
        else:
            out.update(_pdf_facts(path))
    elif suffix in (".md", ".txt", ".html", ".htm") and size <= MAX_TEXT_CHARS:
        out.update(_text_facts(path))
    return out


def inventory(workspace, since=0.0, limit=MAX_FILES, exclude=()):
    """Documents written under `workspace` at/after `since`, oldest first.

    `since` is the turn's start: inputs staged before the turn carry an older
    mtime and are therefore never reported as something the model produced.
    `exclude` names directories that are the HARNESS'S OWN state (data, sessions,
    harness records). They can live inside a workspace, and a record the harness
    wrote about the turn is not a deliverable of the turn — measured 2026-10-30:
    a chat whose workspace held `LONGRUN_*` state dirs had its own harness note
    read back as an artifact.
    """
    root = str(workspace or "").strip()
    if not root or not os.path.isdir(root):
        return []
    skip = {os.path.abspath(os.path.expanduser(str(p))) for p in (exclude or ()) if p}
    if os.path.abspath(root) in skip:
        return []
    found = []
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames
                       if d not in SKIP_DIRS
                       and os.path.abspath(os.path.join(dirpath, d)) not in skip]
        for name in filenames:
            if not name.lower().endswith(DELIVERABLE_EXTS) or name.startswith("."):
                continue
            full = os.path.join(dirpath, name)
            if os.path.abspath(full) in skip:
                continue
            if _mtime(full) + 1.0 < float(since or 0.0):
                continue
            found.append(full)
    found.sort(key=_mtime)
    entries = []
    for full in found[-int(limit):]:
        try:
            rel = os.path.relpath(full, root)
        except ValueError:
            rel = full
        entries.append({"path": rel, "kind": Path(full).suffix.lower().lstrip("."),
                        "facts": facts(full)})
    return entries


#: A source a task can hand over: documents PLUS the structured data formats that
#: are never deliverables but are always evidence.
INPUT_EXTS = DELIVERABLE_EXTS + (".json", ".xml", ".yaml", ".yml", ".db", ".sqlite")

#: The formats a figure comes FROM. A prose brief is an input, not a source: it
#: tells you what to do, it does not supply the numbers, so a deliverable is not
#: required to cite it.
DATA_EXTS = (".csv", ".tsv", ".xlsx", ".xlsm", ".json", ".xml", ".yaml", ".yml",
             ".db", ".sqlite")


def inputs(workspace, since=0.0, limit=12, exclude=()):
    """The files the turn was GIVEN: they existed before it started.

    The mirror image of `inventory` — same workspace, opposite side of `since`. A
    task normally hands the agent a set of sources; an answer that never accounts
    for one of them rests on an input nobody looked at. That is a fact the harness
    can state without knowing the task, so it states it.
    """
    root = str(workspace or "").strip()
    if not root or not os.path.isdir(root):
        return []
    skip = {os.path.abspath(os.path.expanduser(str(p))) for p in (exclude or ()) if p}
    if os.path.abspath(root) in skip:
        return []
    found = []
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames
                       if d not in SKIP_DIRS
                       and os.path.abspath(os.path.join(dirpath, d)) not in skip]
        for name in filenames:
            if not name.lower().endswith(INPUT_EXTS) or name.startswith("."):
                continue
            full = os.path.join(dirpath, name)
            if os.path.abspath(full) in skip:
                continue
            if _mtime(full) + 1.0 >= float(since or 0.0):
                continue  # written BY this turn, not given to it
            try:
                found.append(os.path.relpath(full, root))
            except ValueError:
                found.append(full)
    return sorted(found)[:int(limit)]


#: An identifier a RECORD carries: `AP10260`, `CHK-1402`, `EFT-2210`,
#: `REF-2026-0207`, `IC-2025-0115-013`, `AET-5567201`, `M84.451A`, `STA27`.
#: Every token must also contain a DIGIT — checked in `coverage()`, not in the
#: pattern, because `GAS-ODOR` and `DUPLICATE_RISK` are hyphenated words, not
#: records, and they matched the hyphenated alternative (measured on the real
#: dispatch workspace, 2026-10-31).
_ID_RE = re.compile(
    r"\b(?:"
    r"[A-Z]{2,}\d{2,}"                        # AP10260
    r"|[A-Z]{2,}(?:[-_.][A-Z0-9]+)+"          # CHK-1402, REF-2026-0207, AET-5567201
    r"|[A-Z]\d{2}\.[0-9A-Z]{1,4}"             # M84.451A (ICD-10 style)
    r")\b")

#: A file is a RECORD SET only when it carries at least this many identifiers. Below
#: it there is nothing to enumerate, and the fact would be noise.
COVERAGE_MIN_IDS = 3
#: A FAMILY larger than this is a register (`AP` x800 in the AP subledger): the
#: deliverable cannot enumerate it and a demand to would be noise. A family the
#: deliverable already ENGAGED with is the useful unit — it cites CHK-1401 and
#: CHK-1403, so CHK-1402 missing is a fact about its own work.
COVERAGE_MAX_FAMILY = 50
#: How many missing identifiers are named in one line (the rest are counted).
COVERAGE_SHOW = 6
#: How many families are reported per file.
COVERAGE_MAX_FAMILIES = 4
#: How many input files are inspected per read-back (bounded: a review may not walk
#: a 700-file tree on every turn).
COVERAGE_MAX_FILES = 6

_STEM_RE = re.compile(r"[A-Za-z]+")


def _stem(tok):
    """The alphabetic stem of an identifier — `CHK-1402` -> `CHK`, `AP10260` -> `AP`."""
    m = _STEM_RE.match(tok)
    return (m.group(0) if m else "")[:6]


def coverage(workspace, since=0.0, limit=MAX_FILES, exclude=()):
    """Identifiers the INPUTS hold that the deliverables never mention.

    The read-back described what was WRITTEN and which input was NAMED; it never
    asked whether the RECORDS inside those inputs were accounted for. Measured
    2026-10-31 (JobBench, same model, fresh official judge on the same artefacts):
    bookkeeping longrun 26/88 against openclaw 52/88, medical longrun 12/78 against
    opencode 58/78. Every rubric lost was "state X for THIS record" — `AP10260`
    (visible to the turn, absent from its own vendor table), `CHK-1402` unflagged,
    a per-issue corrective action nobody wrote. A summary that never names the
    record looks complete and is not.

    The unit is the FAMILY the deliverable already engaged with, not the whole file:
    a report that cites `CHK-1401` and `CHK-1403..1406` is told that `CHK-1402` of
    the same six is absent, while the 800-row `AP` register (a family the report
    never enumerates and could not) is not held against it.

    Returns `[{"path", "total", "families": [{"stem", "total", "missing": [...]}]}]`,
    at most one entry per input record set, facts only. Never raises.
    """
    try:
        given = inputs(workspace, since=since, exclude=exclude)
        if not given:
            return []
        entries = inventory(workspace, since=since, limit=limit, exclude=exclude)
        if not entries:
            return []
        parts, used = [], 0
        for entry in entries:
            text = deliverable_text(os.path.join(workspace, entry["path"]))
            if text:
                parts.append(text)
                used += len(text)
                if used >= MAX_TEXT_CHARS:
                    break
        blob = " ".join(parts)[:MAX_TEXT_CHARS]
        if not blob:
            return []
        out = []
        for name in given[:COVERAGE_MAX_FILES]:
            if os.path.splitext(name)[1].lower() not in DATA_EXTS:
                continue  # a brief or a README is not a record set
            text = deliverable_text(os.path.join(workspace, name))
            if not text:
                continue  # .db and other unreadable-in-place formats: no claim
            ids, seen = [], set()
            for m in _ID_RE.finditer(text):
                tok = m.group(0)
                if tok in seen or len(tok) < 4:
                    continue
                if not any(ch.isdigit() for ch in tok):
                    continue  # `GAS-ODOR` is a hyphenated word, not a record
                seen.add(tok)
                ids.append(tok)
            if len(ids) < COVERAGE_MIN_IDS:
                continue
            families = {}
            for tok in ids:
                families.setdefault(_stem(tok), []).append(tok)
            gaps = []
            for stem, members in families.items():
                if not stem or len(members) > COVERAGE_MAX_FAMILY:
                    continue  # a register, not a checklist
                if not any(t in blob for t in members):
                    continue  # the deliverable never engaged this family at all
                missing = [t for t in members if t not in blob]
                if missing:
                    gaps.append({"stem": stem, "total": len(members), "missing": missing})
            if gaps:
                out.append({"path": name, "total": len(ids), "families": gaps})
        return out
    except Exception:  # noqa: BLE001 — the audit may never break a turn
        return []


def _coverage_bits(item):
    """`CHK: 1 of 6 never appear (CHK-1402)` — the families, self-contained."""
    bits = []
    for fam in item["families"][:COVERAGE_MAX_FAMILIES]:
        shown = fam["missing"][:COVERAGE_SHOW]
        more = "" if len(fam["missing"]) <= COVERAGE_SHOW else ", …"
        bits.append("%s: %d of %d never appear (%s%s)"
                    % (fam["stem"], len(fam["missing"]), fam["total"],
                       ", ".join(shown), more))
    return "; ".join(bits)


def _coverage_line(item):
    """One line per file: which record FAMILIES the answer left incomplete."""
    return "%s — %s" % (item["path"], _coverage_bits(item))


def _line(entry):
    """One human line per artifact: `path — <facts>`."""
    f = entry.get("facts") or {}
    kind = entry.get("kind") or "file"
    if f.get("error"):
        return "%s (%s) — %s" % (entry["path"], kind, f["error"])
    if "sheets" in f:
        def _sheet_cell(s):
            """`Name 20x30` is the USED range; add the FILLED one when it differs."""
            used = "%s %dx%d" % (s["name"], s["rows"], s["cols"])
            filled = s.get("filled_rows")
            if filled is None or filled >= s["rows"]:
                return used
            return "%s [%d row(s) filled]" % (used, filled)
        cells = ", ".join(_sheet_cell(s) for s in f["sheets"][:8])
        hollow = [s["name"] for s in f["sheets"]
                  if (s.get("filled_rows") is not None
                      and s.get("filled_rows") <= 1 and (s["rows"] or 0) > 1)]
        return ("%s (%s) — %d sheet(s): %s; charts=%d; conditional-formatting=%d%s%s"
                % (entry["path"], kind, len(f["sheets"]), cells, f.get("charts", 0),
                   f.get("conditional_formatting", 0),
                   "; EMPTY" if f.get("size") == 0 else "",
                   ("; NO DATA in " + ", ".join(hollow[:4])) if hollow else ""))
    if "columns" in f:
        return ("%s (csv) — %d row(s), columns: %s"
                % (entry["path"], f.get("rows", 0), ", ".join(f["columns"])))
    if "pages" in f:
        return "%s (pdf) — %d page(s), %d chars on page 1" % (
            entry["path"], f.get("pages", 0), f.get("page1_chars", 0))
    if "lines" in f:
        return "%s (%s) — %d line(s), %d chars%s" % (
            entry["path"], kind, f.get("lines", 0), f.get("chars", 0),
            " (EMPTY)" if f.get("empty") else "")
    tail = (" (%s)" % f["skipped"] if f.get("skipped")
            else (" (EMPTY)" if f.get("empty") else ""))
    return "%s (%s) — %d bytes%s" % (entry["path"], kind, f.get("size", 0), tail)


#: Words that make a VISUAL element part of the REQUEST. A workbook with `charts=0`
#: is not a style preference when the job asked for a picture: it is a missing
#: deliverable. Measured 2026-10-31 (JobBench medical_secretaries): the read-back
#: printed `conditional-formatting=0`, the model answered "done", and the rubric
#: that weights red/yellow highlighting scored 0. The fact was shown and never
#: required.
#: Nouns that make an ARTIFACT part of the request. A turn that writes no file at
#: all when the job asked for a workbook/report/letter used to be told NOTHING: the
#: read-back returned None (nothing to describe) and the turn closed. Measured
#: 2026-10-31 on JobBench medical_secretaries, session `5757c9eb681c`: two failed
#: generator scripts, then a prose answer with `output/` empty and no harness word
#: about it.
_ARTIFACT_WORDS = ("file", "files", "document", "report", "workbook", "spreadsheet",
                   "sheet", "csv", "xlsx", "xls", "docx", "doc", "letter", "memo",
                   "template", "invoice", "presentation", "slide", "pdf", "save the")


def names_artifact(request):
    """True when the request asks for a FILE (so writing nothing is a failure)."""
    text = " ".join(str(request or "").lower().split())
    return any(w in text for w in _ARTIFACT_WORDS)


def nothing_written(request):
    """The read-back for a turn that wrote NOTHING while the request named a file."""
    if not names_artifact(request):
        return None
    return ("READ-BACK (harness: mechanical file inspection): you have written NO "
            "file for this request — the workspace has no new document at all. "
            "The request names a deliverable file, so a prose answer is not the "
            "result. Produce the file now (and if a previous attempt failed, fix "
            "the failure instead of summarising it), or state exactly what blocks "
            "you and what the user must supply.")
_VISUAL_WORDS = ("chart", "graph", "plot", "visual", "histogram", "dashboard")
_FORMAT_WORDS = ("conditional format", "conditional-format", "highlight", "colour-code",
                 "color-code", "red fill", "yellow", "fill color", "fill colour",
                 "color code", "colour code")


def _request_gaps(request, entries):
    """Gaps that only exist because the REQUEST named them (visual / formatting)."""
    text = " ".join(str(request or "").lower().split())
    if not text:
        return []
    books = [e for e in entries if "sheets" in (e.get("facts") or {})]
    if not books:
        return []
    out = []
    if any(w in text for w in _VISUAL_WORDS) and not any(
            (e.get("facts") or {}).get("charts") for e in books):
        out.append("the request names a chart/visual and no workbook has one "
                   "(charts=0): a picture the reader can see is part of the job")
    if any(w in text for w in _FORMAT_WORDS) and not any(
            (e.get("facts") or {}).get("conditional_formatting") for e in books):
        out.append("the request names conditional formatting/highlighting and every "
                   "workbook reports conditional-formatting=0")
    return out


def gaps(workspace, request="", since=0.0, limit=MAX_FILES, exclude=()):
    """Mechanical reasons a delivery is INCOMPLETE — facts, never prose.

    Used to decide whether a turn EARNED one more read-back cycle: a sheet whose
    content stops at the header, a csv with no data rows, an unreadable file, and
    the two request-driven ones above. The strings are only for the log; the
    decision is `bool(gaps(...))`, so boilerplate wording in the checklist can never
    trigger a cycle by accident.
    """
    out = []
    entries = inventory(workspace, since=since, limit=limit, exclude=exclude)
    for entry in entries:
        facts = entry.get("facts") or {}
        if facts.get("error") or facts.get("empty") or facts.get("size") == 0:
            out.append("%s: empty or unreadable" % entry["path"])
            continue
        for sheet in facts.get("sheets") or []:
            filled = sheet.get("filled_rows")
            if filled is not None and filled <= 1 and (sheet.get("rows") or 0) > 1:
                out.append("%s[%s]: content stops at the header (%d filled row(s))"
                           % (entry["path"], sheet["name"], filled))
        if "columns" in facts and not facts.get("rows"):
            out.append("%s: csv with no data rows" % entry["path"])
    out.extend(_request_gaps(request, entries))
    # A record set the inputs hold and the delivery never mentions is incomplete
    # for the same reason an empty sheet is: the reader was handed the source and
    # asked for the records in it. Bounded by `coverage()`'s own caps.
    for item in coverage(workspace, since=since, limit=limit, exclude=exclude):
        out.append("%s: incomplete record family (%s)"
                   % (item["path"], _coverage_bits(item)))
    return out


def review(workspace, request="", since=0.0, limit=MAX_FILES, exclude=()):
    """The read-back observation, or None when this turn wrote no document.

    The wording is deliberately a CHECKLIST, not a grade: the agent must be able to
    point at each requirement in the read-back, and the last line says what to do
    when it cannot.
    """
    entries = inventory(workspace, since=since, limit=limit, exclude=exclude)
    if not entries:
        return None
    lines = ["READ-BACK of what you WROTE this turn (harness: mechanical file "
             "inspection, not a judgement of your work):", ""]
    lines += ["  - " + _line(e) for e in entries]
    lines += ["", "Before you answer, check this against the request:",
              "  * every file/sheet/column/figure the request names must be visible above;",
              "  * an empty artifact, a sheet with no rows, or a requested visual "
              "element that shows `charts=0` is NOT delivered;",
              "  * `Name 20x30` is the USED range and grows with formatting: when a "
              "sheet also prints `[N row(s) filled]` or `NO DATA`, N is what a reader "
              "actually sees. Formatting is not content;",
              "  * if a requirement is missing, fix the artifact now — a requirement "
              "you cannot point at in the read-back is not done."]
    given = inputs(workspace, since=since, exclude=exclude)
    if given:
        lines += ["", "INPUTS you were GIVEN (they existed before this turn): %d file(s) — %s"
                  % (len(given), ", ".join(given)),
                  "  * every input above must be used or explicitly dismissed; a source "
                  "you never opened cannot back a claim."]
        # EV-37: a figure whose source is never NAMED cannot be credited by a
        # reader — or a grader — however correct it is. Mechanical test: is the
        # input's filename inside the text of what you wrote?
        try:
            blob = " ".join(deliverable_text(os.path.join(workspace, e["path"]))
                            for e in entries)[:MAX_TEXT_CHARS]
        except Exception:  # noqa: BLE001
            blob = ""
        if blob:
            # A DATA source must be named; the prose brief must not. Requiring
            # `TASK_INSTRUCTIONS.txt` in the report is noise the model then spends
            # turns on (measured 2026-10-30), and a brief is not a source of
            # figures — the numbers come from the data files.
            unnamed = [g for g in given
                       if os.path.splitext(g)[1].lower() in DATA_EXTS
                       and os.path.basename(g) not in blob]
            if unnamed:
                lines += ["  * NOT NAMED in any deliverable: %s — name the file a "
                          "figure came from; an invisible source cannot be credited."
                          % ", ".join(unnamed[:8])]
        # The CONTENT, not only the name: a source is accounted for when its
        # records are, and "I analysed it" cannot be seen by a reader. Every rubric
        # this harness lost on 2026-10-31 was of this shape ("state X for THIS
        # record"), never "write a nicer report".
        for item in coverage(workspace, since=since, limit=limit, exclude=exclude):
            lines += ["  * RECORDS NOT ACCOUNTED FOR — " + _coverage_line(item)
                      + ". Name each one in the deliverable (a row, a line, a cell) "
                        "or say explicitly why it is out of scope; a summary that "
                        "never mentions it does not cover it."]
    if any((e.get("facts") or {}).get("charts") for e in entries):
        lines += ["  * a chart is a PICTURE: a reader that extracts text sees only your "
                  "cells — the plotted numbers must also appear as cells (a caption "
                  "naming the chart is not the data)."]
    # The request can DEMAND a visual or a formatting rule; then the fact above is
    # not a note but a missing deliverable, and the model is told so by name.
    for _g in _request_gaps(request, entries):
        lines += ["  * " + _g]
    head = " ".join(str(request or "").split())[:600]
    if head:
        lines += ["", "Request (head): " + head]
    return "\n".join(lines)
